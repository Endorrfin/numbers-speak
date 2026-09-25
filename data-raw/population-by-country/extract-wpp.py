#!/usr/bin/env python3
"""extract-wpp.py — WPP 2024 COMPACT workbook → wpp2024-population-2025.csv (the rows prep.ts reads).

Usage: python3 data-raw/population-by-country/extract-wpp.py <path to WPP2024_GEN_F01_DEMOGRAPHIC_INDICATORS_COMPACT.xlsx>

The workbook (26 MB) is not committed; its URL and sha256 are in README.md. This script keeps only what the
entry needs: sheet "Medium variant", Year 2025, rows of Type "Country/Area" plus the "World" row; columns
name, WPP notes, M49 location code, ISO3, ISO2, total population on 1 July (thousands, 3 decimals = persons)
and WPP's own density (persons per km², kept only to cross-check our density in tests — never shipped).
Footnotes referenced by the kept rows are appended from sheet "NOTES" to wpp2024-notes.csv.
"""
import csv, sys, re
from pathlib import Path
import openpyxl

YEAR = 2025
here = Path(__file__).resolve().parent
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
ws = wb['Medium variant']
header = None
out = []
for row in ws.iter_rows(values_only=True):
    if header is None:
        if row[0] == 'Index':
            header = list(row)
            col = {h: i for i, h in enumerate(header)}
        continue
    if row[col['Year']] != YEAR or row[col['Type']] not in ('Country/Area', 'World'):
        continue
    out.append({
        'name': row[col['Region, subregion, country or area *']],
        'notes': str(row[col['Notes']] or '').strip(),
        'location_code': row[col['Location code']],
        'iso3': row[col['ISO3 Alpha-code']] or '',
        'iso2': row[col['ISO2 Alpha-code']] or '',
        'type': row[col['Type']],
        'year': YEAR,
        'pop_jul_thousands': row[col['Total Population, as of 1 July (thousands)']],
        'wpp_density': row[col['Population Density, as of 1 July (persons per square km)']],
    })
with open(here / f'wpp2024-population-{YEAR}.csv', 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=list(out[0].keys()), lineterminator='\n')
    w.writeheader()
    w.writerows(out)
used = {n for r in out for n in re.findall(r'\d+', r['notes'])}
notes = []
for (text,) in wb['NOTES'].iter_rows(max_col=1, values_only=True):
    m = re.match(r'^\((\d+)\)\s*(.*)$', text or '', re.S)
    if m and m.group(1) in used:
        notes.append({'note': m.group(1), 'text': m.group(2).strip()})
with open(here / 'wpp2024-notes.csv', 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=['note', 'text'], lineterminator='\n')
    w.writeheader()
    w.writerows(notes)
print(f'{len(out)} rows ({sum(r["type"] == "Country/Area" for r in out)} countries/areas), {len(notes)} notes')
